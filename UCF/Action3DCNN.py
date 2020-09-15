import os
import sys
import shutil
import h5py

import tensorflow as tf
from tensorflow.python.client import device_lib
import numpy as np
from openpyxl import Workbook

import InputData_fold
import Networks_dong


flags = tf.app.flags
FLAGS = flags.FLAGS

flags.DEFINE_string('flag_log_dir', './log/', 'Directory to put log files.')

flags.DEFINE_integer('flag_max_epochs', 50, 'Maximum number of epochs to train.')

flags.DEFINE_integer('flag_batch_size', 16, 'Batch size which must be divided extractly by the size of dataset.')

flags.DEFINE_float('flag_learning_rate', 0.003, 'Learning rate to define the momentum optimizer.')

flags.DEFINE_boolean('flag_aug_online', True, 'Whether to implement online augmentation.')

flags.DEFINE_boolean('flag_optical_flow', True, 'Whether to use optical flow.')

flags.DEFINE_integer('flag_group', None, 'Which group should be set as validation set.')

flags.DEFINE_integer('flag_main_struct', 1, 'Number to choose the network struct: \
		0 denotes original CNN; \
		1 denotes tt+Quan compressed network.')

flags.DEFINE_string('gpu_number', '1', 'gpu_number')



def run_training(b_gpu_enabled = False, str_restore_ckpt = None):
	with tf.Graph().as_default(), tf.device('/cpu:0'):

		print('Begin to get dataset.')
		dict_dataset, dict_mean_std = InputData_fold.get_one_group_dataset(FLAGS.flag_log_dir, str_restore_ckpt, FLAGS.flag_optical_flow, FLAGS.flag_group, FLAGS.flag_batch_size)
		print('Get dataset has done.')

		tfv_global_step = tf.get_variable('var_global_step', [], tf.int32, tf.constant_initializer(0, tf.int32), trainable = False)

		tfv_train_phase = tf.Variable(True, trainable = False, name = 'var_train_phase', dtype = tf.bool, collections = [])

		tfob_variable_averages = tf.train.ExponentialMovingAverage(0.9, name = 'avg_variable')

		tfph_aug_online = tf.placeholder(dtype = tf.int32, shape = [FLAGS.flag_batch_size], name = 'ph_aug_online')
		tfv_aug_online = tf.Variable(tfph_aug_online, trainable = False, collections = [], name = 'var_aug_online')

		if FLAGS.flag_aug_online == False:
			tfv_aug_online = None
		dict_inputs_batches = InputData_fold.construct_batch_part_two_stream(dict_mean_std, FLAGS.flag_batch_size, tfv_aug_online, FLAGS.flag_optical_flow)
		t_cubes_rgb = dict_inputs_batches['batches']['batch_train_cubes_rgb']
		t_cubes_optical = dict_inputs_batches['batches']['batch_train_cubes_optical']
		t_labels = dict_inputs_batches['batches']['batch_train_labels']
		v_cubes_rgb = dict_inputs_batches['batches']['batch_validation_cubes_rgb']
		v_cubes_optical = dict_inputs_batches['batches']['batch_validation_cubes_optical']
		v_labels = dict_inputs_batches['batches']['batch_validation_labels']

		n_decay_steps = int(30.0 * dict_dataset['train']['train_labels'].shape[0] / FLAGS.flag_batch_size)
		f_learning_rate = tf.train.exponential_decay(FLAGS.flag_learning_rate, tfv_global_step, n_decay_steps, 0.1, staircase = True)
		optim = tf.train.MomentumOptimizer(f_learning_rate, 0.9)

		with tf.device('/gpu:1'):

			batch_cubes_rgb, batch_cubes_optical, batch_labels = tf.cond(tfv_train_phase, lambda: (t_cubes_rgb, t_cubes_optical, t_labels), lambda: (v_cubes_rgb, v_cubes_optical, v_labels))


			if FLAGS.flag_main_struct == 0:
				loss_rgb, eval_rgb, inf_rgb = Networks_dong.original_cnn(batch_cubes_rgb, batch_labels, tfv_train_phase, 'rgb')
				loss_optical, eval_optical, inf_optical = Networks_dong.original_cnn(batch_cubes_optical, batch_labels, tfv_train_phase, 'optical')
			elif FLAGS.flag_main_struct == 1:
				loss_rgb, eval_rgb, inf_rgb = Networks_dong.neural_network_struct_tt_quan_conv_fc(batch_cubes_rgb, batch_labels, tfv_train_phase, 'rgb')
				loss_optical, eval_optical, inf_optical = Networks_dong.neural_network_struct_tt_quan_conv_fc(batch_cubes_optical, batch_labels, tfv_train_phase, 'optical')
			else:
				return


			correct_flags = tf.nn.in_top_k((inf_rgb + inf_optical) / 2, batch_labels, 1, name = 'total_eval')
			total_eval = tf.cast(correct_flags, tf.int32)

			grads_rgb = optim.compute_gradients(loss_rgb)
			grads_optical = optim.compute_gradients(loss_optical)

			grads_rgb = [(tf.clip_by_value(grad, -1., 1.), var) for grad, var in grads_rgb if grad is not None]
			grads_optical = [(tf.clip_by_value(grad, -1., 1.), var) for grad, var in grads_optical if grad is not None]

			tfop_apply_gradients_rgb = optim.apply_gradients(grads_rgb, tfv_global_step)
			tfop_apply_gradients_optical = optim.apply_gradients(grads_optical, tfv_global_step)
			with tf.control_dependencies([tfop_apply_gradients_rgb, tfop_apply_gradients_optical]):
				tfop_normalize_gs = tfv_global_step.assign_sub(1)

			tfop_variable_averages_apply = tfob_variable_averages.apply(tf.trainable_variables())

			tfv_train_loss_rgb = tf.Variable(5.0, trainable = False, name = 'var_train_loss', dtype = tf.float32)
			tfv_train_loss_optical = tf.Variable(5.0, trainable = False, name = 'var_train_loss', dtype = tf.float32)
			tfv_train_precision = tf.Variable(0.0, trainable = False, name = 'var_train_precision', dtype = tf.float32)

			l_ops_train_lp_update = []
			l_ops_train_lp_update.append(tfv_train_loss_rgb.assign_sub(0.1 * (tfv_train_loss_rgb - loss_rgb)))
			l_ops_train_lp_update.append(tfv_train_loss_optical.assign_sub(0.1 * (tfv_train_loss_optical - loss_optical)))			
			new_precision = tf.reduce_mean(tf.cast(total_eval, tf.float32))
			l_ops_train_lp_update.append(tfv_train_precision.assign_sub(0.1 * (tfv_train_precision - new_precision)))
			tfop_train_lp_update = tf.group(*l_ops_train_lp_update)

			tfop_train = tf.group(tfop_apply_gradients_rgb, tfop_apply_gradients_optical, tfop_normalize_gs, tfop_variable_averages_apply, tfop_train_lp_update)


		tfob_saver = tf.train.Saver(tf.global_variables())
		tfob_saver_ema = tf.train.Saver(tfob_variable_averages.variables_to_restore())

		if b_gpu_enabled == True:
			tfob_sess = tf.Session(config = tf.ConfigProto(allow_soft_placement = True))
		else:
			tfob_sess = tf.Session(config = tf.ConfigProto(allow_soft_placement = True, device_count = {'GPU': 0}))

		tfob_sess.run(tf.global_variables_initializer())

		tfob_coord = tf.train.Coordinator()
		th_threads = tf.train.start_queue_runners(tfob_sess, tfob_coord)

		n_epoch_steps = int(dict_dataset['train']['train_labels'].shape[0] / FLAGS.flag_batch_size + 0.5)
		n_start_epoch = 0
		if str_restore_ckpt is not None:
			tfob_saver.restore(tfob_sess, str_restore_ckpt)
			sys.stdout.write('Previously started training session restored from "%s".\n' % str_restore_ckpt)
			n_start_epoch = int(tfob_sess.run(tfv_global_step)) // n_epoch_steps
		sys.stdout.write('Starting with epoch #%d.\n' % (n_start_epoch + 1))

		l_rc_loss_pre = []
		if os.path.exists(FLAGS.flag_log_dir + '/learning_curve.h5'):
			with h5py.File(FLAGS.flag_log_dir + '/learning_curve.h5', 'r') as file:
				arr_rc_loss_pre = file.get('curve').value
			l_rc_loss_pre = arr_rc_loss_pre.tolist()

		for n_epoch in range(n_start_epoch, FLAGS.flag_max_epochs):
			sys.stdout.write('\n')

			# -------------------------------------------------------------------------------------------------
			# Training begin! 
			tfob_sess.run(tfv_train_phase.assign(True))
			sys.stdout.write('Epoch #%d. [Train]\n' % (n_epoch + 1))
			sys.stdout.flush()


			n_step = 0
			n_index = 0


			while n_step < n_epoch_steps:

				if FLAGS.flag_aug_online == True:
					arr_mark_aug = np.random.uniform(0, 100, (FLAGS.flag_batch_size))
					tfob_sess.run(tfv_aug_online.initializer, feed_dict = {tfph_aug_online: arr_mark_aug})
				

				dict_input_feed = InputData_fold.get_batch_part_train(dict_dataset, dict_mean_std, dict_inputs_batches['input_placeholders'], n_index, FLAGS.flag_batch_size, FLAGS.flag_optical_flow)
				assert dict_input_feed is not None, 'There are some invalid stacks.'


				_, loss_train_rgb, loss_train_optical, eval_train_rgb, eval_train_optical = tfob_sess.run([tfop_train, loss_rgb, loss_optical, eval_rgb, eval_optical], dict_input_feed)
				assert not np.isnan(loss_train_rgb) and not np.isnan(loss_train_optical), 'Model diverged with loss = NaN.'
				n_step += 1
				n_index += FLAGS.flag_batch_size

				sys.stdout.write('Epoch #%d [Train]. Step %d/%d. Batch rgb loss = %.2f. Batch opticalflow loss = %.2f. Batch rgb precision = %.2f. Batch opticalflow precision = %.2f.' % 
					 (n_epoch + 1, n_step, n_epoch_steps, loss_train_rgb, loss_train_optical, np.mean(eval_train_rgb) * 100.0, np.mean(eval_train_optical) * 100.0))
				sys.stdout.write('\n')
				sys.stdout.flush()

			# Training end!checkpoint
			train_loss_value_rgb, train_loss_value_optical, train_precision_value = tfob_sess.run([tfv_train_loss_rgb, tfv_train_loss_optical, tfv_train_precision])
			sys.stdout.write('Epoch #%d. Train rgb loss = %.2f. Train opticalflow loss = %.2f. Train precision = %.2f.\n' % 
					(n_epoch + 1, train_loss_value_rgb, train_loss_value_optical, train_precision_value * 100.0))
			str_checkpoint_path = os.path.join(FLAGS.flag_log_dir, 'model.ckpt')
			str_ckpt = tfob_saver.save(tfob_sess, str_checkpoint_path, tfv_global_step)
			sys.stdout.write('Checkpoint "%s" is saved.\n\n' % str_ckpt)
			# -------------------------------------------------------------------------------------------------

			# -------------------------------------------------------------------------------------------------
			# Evaluate begin! 
			tfob_sess.run(tfv_train_phase.assign(False))
			sys.stdout.write('Epoch #%d. [Evaluation]\n' % (n_epoch + 1))
			tfob_saver_ema.restore(tfob_sess, str_ckpt)
			sys.stdout.write('EMA variables restored.\n')


			n_val_count = dict_dataset['validation']['validation_labels'].shape[0]
			n_val_steps = (n_val_count + FLAGS.flag_batch_size - 1) // FLAGS.flag_batch_size

			n_index = 0

			n_val_corrects = 0
			n_val_losses_rgb = 0.0
			n_val_losses_optical = 0.0

			while n_val_count > 0:
				dict_input_feed = InputData_fold.get_batch_part_validation(dict_dataset, dict_mean_std, dict_inputs_batches['input_placeholders'], n_index, FLAGS.flag_batch_size, FLAGS.flag_optical_flow)
				assert dict_input_feed is not None, 'There are some invalid stacks.'


				eval_validation, loss_validation_rgb, loss_validation_optical = tfob_sess.run([total_eval, loss_rgb, loss_optical], dict_input_feed)
				n_cnt = min(eval_validation.shape[0], n_val_count)
				n_val_count -= n_cnt
				n_cur_step = n_val_steps - (n_val_count + FLAGS.flag_batch_size - 1) // FLAGS.flag_batch_size
				n_index += FLAGS.flag_batch_size

				n_val_corrects += np.sum(eval_validation[:n_cnt])

				n_val_losses_rgb += loss_validation_rgb * FLAGS.flag_batch_size
				n_val_losses_optical += loss_validation_optical * FLAGS.flag_batch_size

				sys.stdout.write('Epoch #%d [Evaluation]. Step %d/%d. Batch rgb loss = %.2f. Batch opticalflow loss = %.2f. Batch precision = %.2f.' % 
					 (n_epoch + 1, n_cur_step, n_val_steps, loss_validation_rgb, loss_validation_optical, np.mean(eval_validation) * 100.0))
				sys.stdout.write('\n')
				sys.stdout.flush()

			# Evaluate end! 
			validation_precision_value = n_val_corrects / dict_dataset['validation']['validation_labels'].shape[0]
			validation_loss_value_rgb = n_val_losses_rgb / dict_dataset['validation']['validation_labels'].shape[0]
			validation_loss_value_optical = n_val_losses_optical / dict_dataset['validation']['validation_labels'].shape[0]
			sys.stdout.write('Epoch #%d. Validation rgb loss = %.2f. Validation opticalflow loss = %.2f. Validation precision = %.2f.\n' % 
					(n_epoch + 1, validation_loss_value_rgb, validation_loss_value_optical, validation_precision_value * 100.0))
			cur_loss_pre = [validation_loss_value_rgb, validation_loss_value_optical, validation_precision_value * 100.0]
			tfob_saver.restore(tfob_sess, str_ckpt)
			sys.stdout.write('Variables restored.\n\n')
			# -------------------------------------------------------------------------------------------------

			l_rc_loss_pre.append(cur_loss_pre)
			with h5py.File(FLAGS.flag_log_dir + '/learning_curve.h5', 'w') as file:
				file.create_dataset('curve', data = np.array(l_rc_loss_pre, dtype = np.float32))

		# loss precision
		wb = Workbook()
		ws = wb.create_sheet()
		for line in l_rc_loss_pre:
			ws.append(line)
		wb.save(FLAGS.flag_log_dir + 'learning_curve.xlsx')
		wb.close()

		tfob_coord.request_stop()
		tfob_coord.join(th_threads)


def main(_):
	os.environ['CUDA_VISIBLE_DEVICES'] = FLAGS.gpu_number
	b_gpu_enabled = False
	l_devices = device_lib.list_local_devices()
	for i in range(len(l_devices)):
		if l_devices[i].device_type == 'GPU':
			if l_devices[i].memory_limit > 4 * 1024 * 1024 * 1024 :
				b_gpu_enabled = True
				break

	str_last_ckpt = tf.train.latest_checkpoint(FLAGS.flag_log_dir)
	if str_last_ckpt is not None:
		while True:
			sys.stdout.write('Checkpoint "%s" found. Continue last training session?\n' % str_last_ckpt)
			sys.stdout.write('Continue - [c/C]. Restart (all content of log dir will be removed) - [r/R]. Abort - [a/A].\n')
			ans = input().lower()
			if len(ans) == 0:
				continue
			if ans[0] == 'c':
				break
			elif ans[0] == 'r':
				str_last_ckpt = None
				shutil.rmtree(FLAGS.flag_log_dir)
				break
			elif ans[0] == 'a':
				return

	if os.path.exists(FLAGS.flag_log_dir) == False:
		os.mkdir(FLAGS.flag_log_dir)

	run_training(b_gpu_enabled, str_last_ckpt)
	print('Program is finished.')


if __name__ == '__main__':
    tf.app.run()
